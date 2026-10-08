import re
from pathlib import Path
from openpyxl import load_workbook

# 你的仓库根目录
ROOT = Path(r"D:\AnimeSubs_rrk")

def get_ep_label(stem: str) -> str:
    """从文件名提取 EP 标签，例如 透明之夜02中日校对 -> EP02"""
    s = stem.upper()

    if "OP" in s:
        return "OP"
    if "ED" in s:
        return "ED"

    # 匹配前后不是数字的两位数字，避免匹配 1646、1P 这种
    m = re.search(r'(?<!\d)(\d{2})(?!\d)', stem)
    if m:
        return f"EP{m.group(1)}"

    return "EP??"

def clean_cell(v) -> str:
    """清理单元格内容：去掉首尾空白，把换行替换成空格"""
    if v is None:
        return ""
    s = str(v).replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
    return s.strip()

def convert_one(xlsx: Path):
    ep = get_ep_label(xlsx.stem)
    wb = load_workbook(xlsx, data_only=True, read_only=True)

    lines = []

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))

        for idx, row in enumerate(rows):
            if row is None:
                continue

            # 确保至少 3 列
            cells = list(row) + [None] * (3 - len(row))

            ja = clean_cell(cells[0])   # A列 日字
            zh = clean_cell(cells[1])   # B列 中字
            sc = clean_cell(cells[2])   # C列 屏幕字

            # 跳过表头行：第一行且第一列是“日字”
            if idx == 0 and ja in ("日字", "日文", "原文", "日本語", "日语"):
                continue

            # 日文或中文都为空则跳过
            if not ja and not zh:
                continue

            # 如果只有中文没日文（比如纯中文行），也照样输出
            if sc:
                line = f"{ep}：{ja}　|　{zh}　【屏幕字：{sc}】"
            else:
                line = f"{ep}：{ja}　|　{zh}"

            lines.append(line)

        # 多个 sheet 之间空一行
        if lines:
            lines.append("")

    wb.close()

    # 去掉结尾多余空行
    while lines and lines[-1] == "":
        lines.pop()

    out = xlsx.with_suffix(".txt")
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"已生成：{out}（{len([l for l in lines if l])} 行）")

# 主流程
for folder in ROOT.iterdir():
    if not folder.is_dir():
        continue
    if folder.name.startswith("."):
        continue

    xlsx_files = sorted(folder.glob("*.xlsx"))
    if not xlsx_files:
        continue

    print(f"\n>>> 处理文件夹：{folder.name}")
    for xlsx in xlsx_files:
        try:
            convert_one(xlsx)
        except Exception as e:
            print(f"  转换失败：{xlsx.name}  原因：{e}")